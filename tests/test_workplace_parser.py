from __future__ import annotations

from datetime import datetime
from pathlib import Path
import sys
import tempfile
import unittest

from openpyxl import Workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "dashboard"))

from workplace_parser import load_workplace  # noqa: E402


class WorkplaceParserTests(unittest.TestCase):
    def test_ts_wip_is_loaded_and_days_are_calculated_only_when_completed(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            base = Path(temporary)
            workplace = base / "workplace"
            workplace.mkdir()

            workbook = Workbook()
            costs = workbook.active
            costs.title = "Costs"
            costs.append(["Invoice Number", "Amount", "Description", "Date", "Period"])
            costs.append(["10112345678", 3759650, "Workstation Monthly Rental for October 2026", datetime(2026, 10, 1), "Oct 2026"])

            wip = workbook.create_sheet("TS-WIP")
            wip.append(["Workplace task tracker"])
            wip.append([])
            wip.append(["Reported Date", "Closed Date", "Activity", "Status", "Comments"])
            wip.append([datetime(2026, 10, 1), datetime(2026, 10, 5), "Repair meeting room display", "Completed", "Closed after testing"])
            wip.append([datetime(2026, 10, 2), None, "Replace pantry chairs", "In progress", "Vendor visit planned"])
            wip.append([datetime(2026, 10, 3), None, "Review signage", "Not started", ""])

            ignored = workbook.create_sheet("Other-WIP")
            ignored.append(["Reported Date", "Closed Date", "Activity", "Status", "Comments"])
            ignored.append([datetime(2026, 10, 1), datetime(2026, 10, 2), "Must not load", "Completed", "Wrong sheet"])
            workbook.save(workplace / "workplace.xlsx")

            payload = load_workplace(base)

            self.assertEqual(payload["meta"]["wip_record_count"], 3)
            by_activity = {row["activity"]: row for row in payload["wip_records"]}
            self.assertEqual(by_activity["Repair meeting room display"]["days_taken"], 4)
            self.assertIsNone(by_activity["Replace pantry chairs"]["days_taken"])
            self.assertIsNone(by_activity["Review signage"]["days_taken"])
            self.assertEqual(by_activity["Repair meeting room display"]["fiscal_year"], "FY26")
            self.assertNotIn("Must not load", by_activity)


if __name__ == "__main__":
    unittest.main()
