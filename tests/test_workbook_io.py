from __future__ import annotations

from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "dashboard"))

from workbook_io import load_workbook_snapshot  # noqa: E402


class WorkbookSnapshotTests(unittest.TestCase):
    def test_transient_onedrive_permission_error_is_retried(self) -> None:
        temporary_directory = Path(tempfile.mkdtemp(prefix="facilities-workbook-io-test-"))
        path = temporary_directory / "Facilities.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = "ready"
        workbook.save(path)
        workbook.close()
        original_read_bytes = Path.read_bytes
        attempts = 0

        def flaky_read(candidate: Path) -> bytes:
            nonlocal attempts
            attempts += 1
            if attempts == 1:
                raise PermissionError("OneDrive is syncing")
            return original_read_bytes(candidate)

        try:
            with patch.object(Path, "read_bytes", flaky_read), patch("workbook_io.time.sleep"):
                loaded = load_workbook_snapshot(path, read_only=True, data_only=True)
            try:
                self.assertEqual(loaded.active["A1"].value, "ready")
                self.assertEqual(attempts, 2)
            finally:
                loaded.close()
        finally:
            path.unlink(missing_ok=True)
            temporary_directory.rmdir()


if __name__ == "__main__":
    unittest.main()
