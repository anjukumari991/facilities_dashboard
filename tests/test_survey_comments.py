from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
import re
from types import SimpleNamespace
import sys
import tempfile
import threading
import unittest
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.utils.datetime import to_excel


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "dashboard"))

from chat_service import DashboardChatService, _summary_payload  # noqa: E402
from survey_parser import load_surveys  # noqa: E402


SURVEY_HEADERS = [
    "ID",
    "Start time",
    "Completion time",
    "Email",
    "Name",
    "How satisfied are you with food and fresh beverage?",
    "How do you rate the recent workplace enhancements?",
    "Work Areas",
    "Meeting Rooms",
    "Pantry",
    "Washroom",
    "Air quality, lighting, temperature and noise",
    "Reception Area",
    "Kitchen / Breakout Area",
    "Overall experience with Sequoia India facilities",
    "Any other Suggestions or Feedback?",
]


class SurveyCommentParserTests(unittest.TestCase):
    def test_repairs_incorrect_a1_dimension_and_accepts_legacy_questions(self) -> None:
        temporary_directory = Path(tempfile.mkdtemp(prefix="facilities-survey-dimension-test-"))
        surveys_directory = temporary_directory / "surveys"
        surveys_directory.mkdir()
        path = surveys_directory / "survey.xlsx"
        rewritten = surveys_directory / "rewritten.xlsx"
        headers = list(SURVEY_HEADERS)
        headers[5] = "Quality of lunch in the pantry"
        headers[6] = "Please rate your overall satisfaction with the recent upgrades to workplace amenities"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers)
        sheet.append(["id", datetime(2026, 10, 6, 12, 25, 38), datetime(2026, 10, 6, 12, 30), "private@example.com", "Private Person", "Satisfied", 4, "Satisfied", "Satisfied", "Satisfied", "Satisfied", "Satisfied", "Satisfied", "Satisfied", 5, ""])
        workbook.save(path)
        workbook.close()
        with ZipFile(path, "r") as source, ZipFile(rewritten, "w") as target:
            for item in source.infolist():
                content = source.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    content = re.sub(br'<dimension ref="[^"]+"', b'<dimension ref="A1"', content, count=1)
                target.writestr(item, content)
        path.unlink()
        rewritten.rename(path)
        try:
            payload = load_surveys(temporary_directory)
            self.assertEqual(payload["meta"]["file_count"], 1)
            self.assertEqual(payload["meta"]["response_count"], 1)
            self.assertEqual(payload["periods"][0]["month"], "2026-10-01")
            self.assertEqual(payload["periods"][0]["food_satisfaction_rate"], 100.0)
            self.assertEqual(payload["periods"][0]["enhancement_average"], 4.0)
        finally:
            path.unlink(missing_ok=True)
            rewritten.unlink(missing_ok=True)
            surveys_directory.rmdir()
            temporary_directory.rmdir()

    def test_reports_privacy_safe_reason_when_start_times_are_invalid(self) -> None:
        temporary_directory = Path(tempfile.mkdtemp(prefix="facilities-survey-diagnostic-test-"))
        surveys_directory = temporary_directory / "surveys"
        surveys_directory.mkdir()
        path = surveys_directory / "survey.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(SURVEY_HEADERS)
        sheet.append(["id", "not a date", "", "private@example.com", "Private Person", "Satisfied", 4, "Satisfied", "Satisfied", "Satisfied", "Satisfied", "Satisfied", "Satisfied", "Satisfied", 4, ""])
        workbook.save(path)
        workbook.close()
        try:
            payload = load_surveys(temporary_directory)
            self.assertEqual(payload["meta"]["discovered_files"], 1)
            self.assertEqual(payload["meta"]["invalid_files"], 1)
            self.assertEqual(payload["meta"]["invalid_reasons"], {"no_valid_start_times": 1})
            self.assertNotIn("Private Person", json.dumps(payload))
            self.assertNotIn("private@example.com", json.dumps(payload))
        finally:
            path.unlink(missing_ok=True)
            surveys_directory.rmdir()
            temporary_directory.rmdir()

    def test_accepts_short_year_text_and_numeric_excel_start_times(self) -> None:
        temporary_directory = Path(tempfile.mkdtemp(prefix="facilities-survey-date-test-"))
        surveys_directory = temporary_directory / "surveys"
        surveys_directory.mkdir()
        path = surveys_directory / "survey.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(SURVEY_HEADERS)
        base_response = [
            "id",
            None,
            None,
            "private@example.com",
            "Private Person",
            "Satisfied",
            4,
            "Satisfied",
            "Satisfied",
            "Satisfied",
            "Satisfied",
            "Satisfied",
            "Satisfied",
            "Satisfied",
            4,
            "",
        ]
        short_year_response = list(base_response)
        short_year_response[1] = "10/6/26 12:25:38"
        numeric_response = list(base_response)
        numeric_response[0] = "id-2"
        numeric_response[1] = to_excel(datetime(2026, 10, 7, 9, 30))
        sheet.append(short_year_response)
        sheet.append(numeric_response)
        workbook.save(path)
        workbook.close()
        try:
            payload = load_surveys(temporary_directory)
            self.assertEqual(payload["meta"]["file_count"], 1)
            self.assertEqual(payload["meta"]["response_count"], 2)
            self.assertEqual(payload["periods"][0]["month"], "2026-10-01")
        finally:
            path.unlink(missing_ok=True)
            surveys_directory.rmdir()
            temporary_directory.rmdir()

    def test_reads_only_recognised_comments_and_masks_direct_identifiers(self) -> None:
        temporary_directory = Path(tempfile.mkdtemp(prefix="facilities-survey-comment-test-"))
        surveys_directory = temporary_directory / "surveys"
        surveys_directory.mkdir()
        path = surveys_directory / "survey.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(SURVEY_HEADERS)
        sheet.append(
            [
                "SECRET-ID-1",
                "2026-10-01 09:00:00",
                "2026-10-01 09:05:00",
                "person@example.com",
                "Private Person",
                "Satisfied",
                4,
                "Satisfied",
                "Satisfied",
                "Satisfied",
                "Satisfied",
                "Satisfied",
                "Satisfied",
                "Satisfied",
                5,
                "Email me at person@example.com; Anju & Suman are helpful; employee ID A-123.",
            ]
        )
        workbook.save(path)
        workbook.close()
        try:
            payload = load_surveys(temporary_directory)
            serialized = json.dumps(payload)
            comment = payload["periods"][0]["comments"][0]["text"]
            self.assertEqual(payload["periods"][0]["comment_count"], 1)
            self.assertEqual(payload["periods"][0]["commenting_response_count"], 1)
            self.assertIn("[email removed]", comment)
            self.assertIn("[names removed]", comment)
            self.assertIn("[employee identifier removed]", comment)
            self.assertNotIn("Private Person", serialized)
            self.assertNotIn("SECRET-ID-1", serialized)
            self.assertNotIn("person@example.com", serialized)
        finally:
            path.unlink(missing_ok=True)
            surveys_directory.rmdir()
            temporary_directory.rmdir()

    def test_does_not_treat_rating_question_feedback_word_as_comment(self) -> None:
        headers = list(SURVEY_HEADERS)
        headers[5] = "Food and fresh beverage feedback and support"
        headers.pop()
        temporary_directory = Path(tempfile.mkdtemp(prefix="facilities-survey-no-comment-test-"))
        surveys_directory = temporary_directory / "surveys"
        surveys_directory.mkdir()
        path = surveys_directory / "survey.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers)
        sheet.append(["id", "2026-10-01", "2026-10-01", "x@example.com", "Name", "Satisfied", 4, "Satisfied", "Satisfied", "Satisfied", "Satisfied", "Satisfied", "Satisfied", "Satisfied", 4])
        workbook.save(path)
        workbook.close()
        try:
            payload = load_surveys(temporary_directory)
            self.assertEqual(payload["periods"][0]["comments"], [])
        finally:
            path.unlink(missing_ok=True)
            surveys_directory.rmdir()
            temporary_directory.rmdir()


class SurveyCommentSummaryTests(unittest.TestCase):
    def test_gateway_timeout_is_reported_as_request_timeout(self) -> None:
        gateway_timeout = type("InternalServerError", (Exception,), {"status_code": 504})
        result = DashboardChatService._safe_error(gateway_timeout("Gateway Timeout"))
        self.assertEqual(result["code"], "request_timeout")
        self.assertIn("did not finish", result["error"])

    def test_timeout_is_not_reported_as_connection_failure(self) -> None:
        timeout_error = type("APITimeoutError", (Exception,), {})
        result = DashboardChatService._safe_error(timeout_error("timed out"))
        self.assertEqual(result["code"], "request_timeout")
        self.assertIn("reachable", result["error"])

    def test_summary_payload_accepts_json_fences_and_validates_mentions(self) -> None:
        parsed = _summary_payload(
            '```json\n{"appreciations":[{"point":"Healthy food","mentions":3}],"employee_asks":[]}\n```',
            2,
        )
        self.assertEqual(parsed["appreciations"][0]["point"], "Healthy food")
        self.assertIsNone(parsed["appreciations"][0]["mentions"])

    def test_summary_payload_accepts_local_model_preamble_and_key_variants(self) -> None:
        parsed = _summary_payload(
            'Here is the requested result:\n{"Appreciations":["Clean meeting rooms"],'
            '"Employee Asks":[{"point":"More lunch variety","mentions":2}],'
            '"Focus Areas":[],"Suggested Actions":["Review the menu"]}\nDone.',
            5,
        )
        self.assertEqual(parsed["appreciations"][0]["point"], "Clean meeting rooms")
        self.assertEqual(parsed["employee_asks"][0]["point"], "Review the menu")
        self.assertNotIn("focus_areas", parsed)
        self.assertNotIn("suggested_actions", parsed)

    def test_summary_payload_removes_repeated_employee_asks(self) -> None:
        parsed = _summary_payload(
            json.dumps(
                {
                    "appreciations": [],
                    "employee_asks": [
                        {"point": "Provide more healthy lunch options", "mentions": 4},
                        {"point": "More healthy options for lunch", "mentions": 4},
                        {"point": "Add more meeting rooms", "mentions": 2},
                    ],
                }
            ),
            8,
        )
        self.assertEqual(len(parsed["employee_asks"]), 2)
        self.assertEqual(parsed["employee_asks"][0]["point"], "Provide more healthy lunch options")

    def test_summary_payload_moves_requests_out_of_appreciations(self) -> None:
        parsed = _summary_payload(
            json.dumps(
                {
                    "appreciations": [
                        {"point": "More meeting rooms", "mentions": 3},
                        {"point": "Clean and welcoming reception", "mentions": 4},
                    ],
                    "employee_asks": [{"point": "Improve lunch variety", "mentions": 2}],
                }
            ),
            8,
        )
        self.assertEqual([item["point"] for item in parsed["appreciations"]], ["Clean and welcoming reception"])
        self.assertEqual(
            [item["point"] for item in parsed["employee_asks"]],
            ["Improve lunch variety", "More meeting rooms"],
        )

    def test_summary_sends_no_identity_fields_and_redacts_comment_content(self) -> None:
        captured: dict = {}

        class Models:
            def list(self):
                return SimpleNamespace(data=[SimpleNamespace(id="local-model")])

        class Completions:
            def create(self, **request):
                captured.update(request)
                content = json.dumps(
                    {
                        "appreciations": [{"point": "Facilities support", "mentions": 1}],
                        "employee_asks": [],
                    }
                )
                return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=content))])

        client = SimpleNamespace(models=Models(), chat=SimpleNamespace(completions=Completions()))
        settings = SimpleNamespace(load=lambda: {"ai_api_endpoint": "http://127.0.0.1:12345/v1"})
        state = SimpleNamespace(
            lock=threading.RLock(),
            library=SimpleNamespace(settings=settings),
            survey_payload={
                "periods": [
                    {
                        "month": "2026-10-01",
                        "comments": [
                            {"topic": "Feedback", "text": "Anju & Suman helped; contact person@example.com"},
                            {"topic": "Feedback", "text": "More meeting rooms"},
                            {"topic": "Feedback", "text": "Healthy lunch options"},
                            {"topic": "Feedback", "text": "Facilities are good"},
                            {"topic": "Feedback", "text": "Improve pantry seating"},
                        ],
                    }
                ]
            },
        )
        credentials = SimpleNamespace(get_openai_api_token=lambda: "test-token")
        service = DashboardChatService(state, credentials, client_factory=lambda *_args: client)
        result = service.summarize_survey_comments("2026-10-01")
        prompt = captured["messages"][1]["content"]
        self.assertTrue(result["ok"])
        self.assertNotIn("Anju", prompt)
        self.assertNotIn("Suman", prompt)
        self.assertNotIn("person@example.com", prompt)
        self.assertNotIn("email", json.loads(prompt)["comments"][0])

    def test_managed_summary_uses_local_connection_and_constrained_schema(self) -> None:
        captured: dict = {}

        class Models:
            def list(self):
                return SimpleNamespace(data=[SimpleNamespace(id="qwen-local")])

        class Completions:
            def create(self, **request):
                captured["request"] = request
                content = json.dumps(
                    {
                        "appreciations": [{"point": "Clean facilities", "mentions": 3}],
                        "employee_asks": [],
                    }
                )
                return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=content))])

        def client_factory(token, endpoint):
            captured["token"] = token
            captured["endpoint"] = endpoint
            return SimpleNamespace(models=Models(), chat=SimpleNamespace(completions=Completions()))

        settings = SimpleNamespace(load=lambda: {"ai_provider": "managed", "ai_api_endpoint": "https://unused.example/v1"})
        state = SimpleNamespace(
            lock=threading.RLock(),
            library=SimpleNamespace(settings=settings),
            survey_payload={
                "periods": [
                    {
                        "month": "2026-10-01",
                        "comments": [{"topic": "Feedback", "text": f"Anonymous feedback {index}"} for index in range(5)],
                    }
                ]
            },
        )
        credentials = SimpleNamespace(
            get_openai_api_token=lambda: "external-token",
            get_local_ai_token=lambda: "local-token",
        )
        service = DashboardChatService(
            state,
            credentials,
            client_factory=client_factory,
            managed_endpoint=lambda: "http://127.0.0.1:12481/v1",
        )

        result = service.summarize_survey_comments("2026-10-01")

        self.assertTrue(result["ok"])
        self.assertEqual(result["provider"], "managed")
        self.assertEqual(captured["token"], "local-token")
        self.assertEqual(captured["endpoint"], "http://127.0.0.1:12481/v1")
        response_format = captured["request"]["response_format"]
        self.assertEqual(response_format["type"], "json_object")
        self.assertEqual(response_format["schema"]["required"], ["appreciations", "employee_asks"])


if __name__ == "__main__":
    unittest.main()
