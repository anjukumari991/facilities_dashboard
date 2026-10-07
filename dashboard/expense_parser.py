from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import re
from typing import Any, Iterable

from openpyxl import load_workbook


MONTH_FORMATS = (
    "%b'%y",
    "%b %Y",
    "%b %y",
    "%B %Y",
    "%b-%Y",
    "%b-%y",
    "%B-%Y",
    "%Y-%m",
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%m/%Y",
    "%m-%Y",
    "%d/%m/%Y",
    "%d-%m-%Y",
    "%m/%d/%Y",
)


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = re.sub(r"[^0-9.\-]", "", str(value))
    if not cleaned or cleaned in {"-", "."}:
        return None
    try:
        return float(cleaned)
    except ValueError:
        return None


def _month(value: Any, month_year_hints: dict[str, int] | None = None) -> date | None:
    if isinstance(value, datetime):
        return value.date().replace(day=1)
    if isinstance(value, date):
        return value.replace(day=1)
    raw = _text(value)
    if not raw:
        return None
    numeric_date = re.fullmatch(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", raw)
    if numeric_date:
        first, second, year = (int(part) for part in numeric_date.groups())
        try:
            # Expense Month should be the first day of the month. Accept both
            # US-style (9/1/2026) and Indian-style (1/9/2026) Excel displays.
            if second == 1 and 1 <= first <= 12:
                return date(year, first, 1)
            if first == 1 and 1 <= second <= 12:
                return date(year, second, 1)
        except ValueError:
            return None
    for fmt in MONTH_FORMATS:
        try:
            return datetime.strptime(raw, fmt).date().replace(day=1)
        except ValueError:
            pass
    short = raw[:3].title()
    if month_year_hints and short in month_year_hints:
        try:
            return datetime.strptime(f"{short} {month_year_hints[short]}", "%b %Y").date()
        except ValueError:
            return None
    return None


def _clean_vendor_category(label: str) -> tuple[str, str]:
    raw = re.sub(r"\s+", " ", label).strip()
    mappings = [
        (r"^Clientfy\s*-\s*(.+)$", "Clientfy", None),
        (r"^Book\s*Mark(?:\s*-\s*(.+))?$", "Book Mark", "Pantry"),
        (r"^Flowers(?:\s*-\s*(.+))?$", "Flowers", "Flowers"),
        (r"^Taaza Shahi Mewa\s*-\s*(.+)$", "Taaza Shahi Mewa", None),
        (r"^Flavvora\s*-\s*(.+)$", "Flavvora", None),
        (r"^Awfis Space Solutions Limited$", "Awfis Space Solutions Limited", "Workspace"),
        (r"^Ambassador Tours and Travels LLP$", "Ambassador Tours and Travels LLP", "Travel"),
        (r"^Courier$", "Unspecified", "Courier"),
        (r"^Fruit(?:\s*&\s*Snacks)?$", "Unspecified", "Fruit & Snacks"),
    ]
    for pattern, vendor, default_category in mappings:
        match = re.match(pattern, raw, flags=re.IGNORECASE)
        if not match:
            continue
        captured = match.group(1).strip() if match.lastindex and match.group(1) else None
        category = captured or default_category or "Other"
        category = re.sub(r"\bStationary\b", "Stationery", category, flags=re.IGNORECASE)
        return vendor, category
    if "-" in raw:
        vendor, category = [part.strip() for part in raw.split("-", 1)]
        return vendor or "Unspecified", category or "Other"
    return raw or "Unspecified", "Other"


def _report_category(expense_group: str, category: str, vendor: str = "") -> str:
    if expense_group.casefold() == "lunch":
        return "Lunch"
    lowered = category.casefold()
    vendor_lowered = vendor.casefold()
    if "fruit" in lowered:
        return "Fruits"
    if "flower" in vendor_lowered or any(term in lowered for term in ("stationery", "stationary", "flower", "medicine")):
        return "Miscellaneous"
    if any(term in lowered for term in ("pantry", "snack")):
        return "Pantry"
    if "workspace" in lowered:
        return "Workspace"
    if "travel" in lowered:
        return "Travel"
    if "courier" in lowered:
        return "Courier"
    if category and lowered != "other":
        return category
    if vendor and vendor_lowered not in {"other", "unspecified"}:
        return vendor
    return "Unclassified"


def _record_key(record: dict[str, Any]) -> tuple[str, str, str, str]:
    return (
        record["month"],
        record["expense_group"],
        record["vendor"].casefold(),
        record["category"].casefold(),
    )


def _material_difference(old: dict[str, Any], new: dict[str, Any]) -> bool:
    fields = (
        "amount_inr",
        "food_orders",
        "counts_ordered",
        "actual_count",
        "per_person_cost_inr",
    )
    for field in fields:
        left, right = old.get(field), new.get(field)
        if left is None and right is None:
            continue
        if left is None or right is None or abs(float(left) - float(right)) > 0.01:
            return True
    return False


def _dedupe(records: Iterable[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    latest: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    conflicts: list[dict[str, Any]] = []
    for record in records:
        key = _record_key(record)
        previous = latest.get(key)
        if previous and _material_difference(previous, record):
            conflicts.append(
                {
                    "month": record["month"],
                    "expense_group": record["expense_group"],
                    "vendor": record["vendor"],
                    "category": record["category"],
                    "previous_source": previous["source_ref"],
                    "chosen_source": record["source_ref"],
                    "previous_amount_inr": previous.get("amount_inr"),
                    "chosen_amount_inr": record.get("amount_inr"),
                }
            )
        latest[key] = record
    ordered = sorted(latest.values(), key=lambda row: (row["month"], row["expense_group"], row["vendor"], row["category"]))
    return ordered, conflicts


def _parse_normalized(workbook) -> tuple[list[dict[str, Any]], list[dict[str, Any]]] | None:
    if "Expenses" not in workbook.sheetnames:
        return None
    sheet = workbook["Expenses"]
    headers = {_text(cell.value): idx + 1 for idx, cell in enumerate(sheet[1]) if _text(cell.value)}
    required = {"Expense Month", "Expense Group", "Vendor", "Category", "Amount INR"}
    if not required.issubset(headers):
        return None
    rows: list[dict[str, Any]] = []
    for row in range(2, sheet.max_row + 1):
        month_value = sheet.cell(row, headers["Expense Month"]).value
        parsed_month = _month(month_value)
        amount = _number(sheet.cell(row, headers["Amount INR"]).value)
        if not parsed_month or amount is None:
            continue
        def value(name: str) -> Any:
            column = headers.get(name)
            return sheet.cell(row, column).value if column else None
        actual = _number(value("Actual Count"))
        per_person = _number(value("Per Person Cost INR"))
        if per_person is None and actual:
            per_person = amount / actual
        expense_group = _text(value("Expense Group")) or "Other"
        category = _text(value("Category")) or "Other"
        vendor = _text(value("Vendor")) or "Unspecified"
        report_category = _text(value("Report Category"))
        rows.append(
            {
                "month": parsed_month.isoformat(),
                "expense_group": expense_group,
                "report_category": report_category if report_category.casefold() not in {"", "other"} else _report_category(expense_group, category, vendor),
                "vendor": vendor,
                "category": category,
                "amount_inr": amount,
                "food_orders": _number(value("Food Orders")),
                "counts_ordered": _number(value("Counts Ordered")),
                "actual_count": actual,
                "per_person_cost_inr": per_person,
                "source_sheet": "Expenses",
                "source_ref": f"Expenses!{row}:{row}",
                "notes": _text(value("Notes")),
            }
        )
    return _dedupe(rows)


def _parse_reorganized(workbook) -> tuple[list[dict[str, Any]], list[dict[str, Any]]] | None:
    if not {"Lunch", "Pantry & Misc"}.issubset(workbook.sheetnames):
        return None
    lunch_sheet = workbook["Lunch"]
    pantry_sheet = workbook["Pantry & Misc"]
    lunch_headers = {_text(lunch_sheet.cell(4, idx).value): idx for idx in range(1, 14) if _text(lunch_sheet.cell(4, idx).value)}
    pantry_headers = {_text(pantry_sheet.cell(4, idx).value): idx for idx in range(1, 12) if _text(pantry_sheet.cell(4, idx).value)}
    if not {"Expense Month", "Vendor", "Invoice Value INR"}.issubset(lunch_headers):
        return None
    if not {"Expense Month", "Report Category", "Vendor", "Expense Type", "Amount INR"}.issubset(pantry_headers):
        return None

    rows: list[dict[str, Any]] = []
    conflicts: list[dict[str, Any]] = []
    for row in range(5, lunch_sheet.max_row + 1):
        status = _text(lunch_sheet.cell(row, lunch_headers["Version Status"]).value) if "Version Status" in lunch_headers else "Current"
        if status == "Historical - differs":
            conflicts.append({"sheet": "Lunch", "row": row, "status": status})
        if "Current Record" in lunch_headers and _text(lunch_sheet.cell(row, lunch_headers["Current Record"]).value).casefold() != "yes":
            continue
        parsed_month = _month(lunch_sheet.cell(row, lunch_headers["Expense Month"]).value)
        amount = _number(lunch_sheet.cell(row, lunch_headers["Invoice Value INR"]).value)
        if not parsed_month or amount is None:
            continue
        actual = _number(lunch_sheet.cell(row, lunch_headers["Actual Count"]).value)
        per_person = _number(lunch_sheet.cell(row, lunch_headers["Per Person Cost INR"]).value)
        rows.append({
            "month": parsed_month.isoformat(),
            "expense_group": "Lunch",
            "report_category": "Lunch",
            "vendor": _text(lunch_sheet.cell(row, lunch_headers["Vendor"]).value),
            "category": "Employee Meals",
            "amount_inr": amount,
            "food_orders": _number(lunch_sheet.cell(row, lunch_headers["Food Orders"]).value),
            "counts_ordered": _number(lunch_sheet.cell(row, lunch_headers["Counts Ordered"]).value),
            "actual_count": actual,
            "per_person_cost_inr": per_person if per_person is not None else (amount / actual if actual else None),
            "source_sheet": "Lunch",
            "source_ref": (_text(lunch_sheet.cell(row, lunch_headers["Source Cell"]).value) if "Source Cell" in lunch_headers else "") or f"Lunch!{row}:{row}",
            "notes": status,
        })
    for row in range(5, pantry_sheet.max_row + 1):
        status = _text(pantry_sheet.cell(row, pantry_headers["Version Status"]).value) if "Version Status" in pantry_headers else "Current"
        if status == "Historical - differs":
            conflicts.append({"sheet": "Pantry & Misc", "row": row, "status": status})
        if "Current Record" in pantry_headers and _text(pantry_sheet.cell(row, pantry_headers["Current Record"]).value).casefold() != "yes":
            continue
        parsed_month = _month(pantry_sheet.cell(row, pantry_headers["Expense Month"]).value)
        amount = _number(pantry_sheet.cell(row, pantry_headers["Amount INR"]).value)
        if not parsed_month or amount is None:
            continue
        rows.append({
            "month": parsed_month.isoformat(),
            "expense_group": "Pantry & Misc",
            "report_category": _text(pantry_sheet.cell(row, pantry_headers["Report Category"]).value) or "Other",
            "vendor": _text(pantry_sheet.cell(row, pantry_headers["Vendor"]).value) or "Unspecified",
            "category": _text(pantry_sheet.cell(row, pantry_headers["Expense Type"]).value) or "Other",
            "amount_inr": amount,
            "food_orders": None,
            "counts_ordered": None,
            "actual_count": None,
            "per_person_cost_inr": None,
            "source_sheet": "Pantry & Misc",
            "source_ref": (_text(pantry_sheet.cell(row, pantry_headers["Source Cell"]).value) if "Source Cell" in pantry_headers else "") or f"Pantry & Misc!{row}:{row}",
            "notes": status,
        })
    return rows, conflicts


def _parse_lunch(sheet) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for start_col in (2, 10):
        row = 1
        while row <= sheet.max_row:
            if _text(sheet.cell(row, start_col + 1).value).casefold() != "month":
                row += 1
                continue
            vendor = ""
            cursor = row + 1
            while cursor <= min(sheet.max_row, row + 8):
                month_raw = _text(sheet.cell(cursor, start_col + 1).value)
                if not month_raw or month_raw.casefold().startswith("average") or month_raw.casefold() == "month":
                    break
                parsed_month = _month(month_raw)
                if not parsed_month:
                    break
                current_vendor = _text(sheet.cell(cursor, start_col).value)
                if current_vendor:
                    vendor = current_vendor
                amount = _number(sheet.cell(cursor, start_col + 5).value)
                actual_count = _number(sheet.cell(cursor, start_col + 4).value)
                if vendor and amount is not None:
                    per_person = _number(sheet.cell(cursor, start_col + 6).value)
                    if per_person is None and actual_count:
                        per_person = amount / actual_count
                    rows.append(
                        {
                            "month": parsed_month.isoformat(),
                            "expense_group": "Lunch",
                            "vendor": vendor,
                            "category": "Employee Meals",
                            "amount_inr": amount,
                            "food_orders": _number(sheet.cell(cursor, start_col + 2).value),
                            "counts_ordered": _number(sheet.cell(cursor, start_col + 3).value),
                            "actual_count": actual_count,
                            "per_person_cost_inr": per_person,
                            "source_sheet": sheet.title,
                            "source_ref": f"{sheet.title}!{sheet.cell(cursor, start_col).coordinate}:{sheet.cell(cursor, start_col + 6).coordinate}",
                            "notes": "Latest repeated occurrence retained during normalization.",
                        }
                    )
                cursor += 1
            row = max(row + 1, cursor)
    return rows


def _parse_pantry(sheet, month_year_hints: dict[str, int]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for header_row in range(1, sheet.max_row + 1):
        if _text(sheet.cell(header_row, 2).value).casefold() != "facility category":
            continue
        months = [_month(sheet.cell(header_row, col).value, month_year_hints) for col in range(3, 6)]
        row = header_row + 1
        while row <= sheet.max_row:
            label = _text(sheet.cell(row, 2).value)
            if not label or label.casefold().startswith("total monthly") or label.casefold() == "facility category":
                break
            vendor, category = _clean_vendor_category(label)
            for offset, col in enumerate(range(3, 6)):
                parsed_month = months[offset]
                amount = _number(sheet.cell(row, col).value)
                if parsed_month and amount is not None:
                    rows.append(
                        {
                            "month": parsed_month.isoformat(),
                            "expense_group": "Pantry & Misc",
                            "vendor": vendor,
                            "category": category,
                            "amount_inr": amount,
                            "food_orders": None,
                            "counts_ordered": None,
                            "actual_count": None,
                            "per_person_cost_inr": None,
                            "source_sheet": sheet.title,
                            "source_ref": f"{sheet.title}!{sheet.cell(row, col).coordinate}",
                            "notes": f"Normalized from legacy label: {label}",
                        }
                    )
            row += 1
    return rows


def _finalize(records: list[dict[str, Any]], conflicts: list[dict[str, Any]], source_format: str) -> dict[str, Any]:
    for record in records:
        current = _text(record.get("report_category"))
        if current.casefold() in {"", "other"}:
            record["report_category"] = _report_category(record["expense_group"], record["category"], record["vendor"])

    months = sorted({record["month"] for record in records})
    groups = sorted({record["expense_group"] for record in records})
    report_categories = sorted({record["report_category"] for record in records})
    vendors = sorted({record["vendor"] for record in records})
    month_group_counts: dict[str, set[str]] = defaultdict(set)
    for record in records:
        month_group_counts[record["expense_group"]].add(record["month"])

    missing_months: list[dict[str, str]] = []
    if months:
        start = date.fromisoformat(months[0])
        end = date.fromisoformat(months[-1])
        cursor = start
        expected: list[str] = []
        while cursor <= end:
            expected.append(cursor.isoformat())
            cursor = date(cursor.year + (cursor.month == 12), 1 if cursor.month == 12 else cursor.month + 1, 1)
        for group in groups:
            for month_value in expected:
                if month_value not in month_group_counts[group]:
                    missing_months.append({"expense_group": group, "month": month_value})

    return {
        "records": records,
        "meta": {
            "source_format": source_format,
            "record_count": len(records),
            "months": months,
            "groups": groups,
            "report_categories": report_categories,
            "vendors": vendors,
            "conflict_count": len(conflicts),
            "conflicts": conflicts,
            "missing_months": missing_months,
        },
    }


def load_expenses(source: str | Path | bytes | BytesIO) -> dict[str, Any]:
    workbook = load_workbook(source, data_only=True, read_only=False)
    normalized = _parse_normalized(workbook)
    if normalized is not None:
        records, conflicts = normalized
        return _finalize(records, conflicts, "normalized")

    reorganized = _parse_reorganized(workbook)
    if reorganized is not None:
        records, conflicts = reorganized
        return _finalize(records, conflicts, "reorganized")

    lunch_rows = _parse_lunch(workbook["Lunch"]) if "Lunch" in workbook.sheetnames else []
    hints: dict[str, int] = {}
    for record in lunch_rows:
        parsed = date.fromisoformat(record["month"])
        hints[parsed.strftime("%b")] = parsed.year
    pantry_rows = (
        _parse_pantry(workbook["Pantry & Misc"], hints)
        if "Pantry & Misc" in workbook.sheetnames
        else []
    )
    records, conflicts = _dedupe([*lunch_rows, *pantry_rows])
    return _finalize(records, conflicts, "legacy")
