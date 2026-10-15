from __future__ import annotations

from io import BytesIO
from pathlib import Path
import time
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


class WorkbookTemporarilyUnavailable(PermissionError):
    """A workbook could not be snapshotted while Excel/OneDrive was using it."""


def load_workbook_snapshot(source: Any, **options: Any) -> Any:
    """Open a local workbook from a short-lived, retrying byte snapshot.

    Reading into memory releases the OneDrive/Excel file handle immediately.
    It also prevents openpyxl's lazy ZIP reads from racing a cloud-sync update.
    """
    if isinstance(source, bytes):
        return load_workbook(BytesIO(source), **options)
    if isinstance(source, BytesIO) or not isinstance(source, (str, Path)):
        return load_workbook(source, **options)

    path = Path(source)
    delays = (0.2, 0.5, 1.0, 2.0)
    last_error: Exception | None = None
    for attempt in range(len(delays) + 1):
        try:
            before = path.stat()
            content = path.read_bytes()
            after = path.stat()
            if not content or len(content) != after.st_size:
                raise OSError("The workbook changed while it was being read.")
            if before.st_mtime_ns != after.st_mtime_ns or before.st_size != after.st_size:
                raise OSError("The workbook changed while it was being read.")
            return load_workbook(BytesIO(content), **options)
        except (PermissionError, FileNotFoundError, OSError, BadZipFile, InvalidFileException, EOFError) as exc:
            last_error = exc
            if attempt < len(delays):
                time.sleep(delays[attempt])

    raise WorkbookTemporarilyUnavailable(
        f"Could not read '{path.name}'. Excel or OneDrive may be locking or syncing it. "
        "Close the workbook in Excel, wait for the OneDrive green check mark, or choose "
        "'Always keep on this device', then select Refresh."
    ) from last_error
