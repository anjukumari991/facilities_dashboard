from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
import hashlib
from pathlib import Path
import re
from typing import Any, Iterable

from openpyxl import load_workbook


SUPPORTED_WORKBOOK_SUFFIXES = {".xlsx", ".xlsm"}
REQUIRED_COLUMNS = ("Invoice Number", "Amount", "Description", "Date", "Period")
CONTRACT_ICON_TYPES = {
    ".pdf": "pdf",
    ".doc": "word",
    ".docx": "word",
    ".xls": "excel",
    ".xlsx": "excel",
    ".xlsm": "excel",
    ".ppt": "presentation",
    ".pptx": "presentation",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".txt": "text",
}


def _normalized(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def _header_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", _normalized(value))


HEADER_KEYS = {_header_key(name): name for name in REQUIRED_COLUMNS}


def _amount(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        amount = float(value)
    else:
        raw = re.sub(r"[^0-9.()\-]", "", str(value or ""))
        if not raw:
            return None
        if raw.startswith("(") and raw.endswith(")"):
            raw = f"-{raw[1:-1]}"
        try:
            amount = float(raw)
        except ValueError:
            return None
    return round(amount, 2) if amount >= 0 else None


def _period_month(value: Any) -> date | None:
    raw = re.sub(r"\s+", " ", str(value or "")).strip()
    for fmt in ("%b %Y", "%B %Y", "%Y-%m", "%m/%Y"):
        try:
            parsed = datetime.strptime(raw, fmt)
            return date(parsed.year, parsed.month, 1)
        except ValueError:
            pass
    return None


def _invoice_date(value: Any, period: Any = None) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value or "").strip()
    candidates: list[date] = []
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%d-%m-%Y", "%m-%d-%Y"):
        try:
            candidates.append(datetime.strptime(raw, fmt).date())
        except ValueError:
            pass
    expected = _period_month(period)
    if expected:
        matching = [candidate for candidate in candidates if (candidate.year, candidate.month) == (expected.year, expected.month)]
        if matching:
            return matching[0]
    return candidates[0] if candidates else None


def _invoice_number(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value or "").strip()


def fiscal_year_for_month(month: date) -> str:
    start_year = month.year if month.month >= 5 else month.year - 1
    return f"FY{start_year % 100:02d}"


def _cost_type(description: str) -> str:
    normalized = _normalized(description)
    if "monthly rental" in normalized or ("workstation" in normalized and "rent" in normalized):
        return "Monthly rental"
    cleaned = re.sub(
        r"\b(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|jul(?:y)?|aug(?:ust)?|sep(?:tember)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)\s+20\d{2}\b",
        "",
        description,
        flags=re.IGNORECASE,
    )
    return re.sub(r"\s+", " ", cleaned).strip(" -–—,:;")[:120] or "Other workplace cost"


def _workplace_sheets(workbook) -> Iterable[tuple[Any, int, dict[str, int]]]:
    for sheet in workbook.worksheets:
        for row_number in range(1, min(sheet.max_row, 15) + 1):
            mapping: dict[str, int] = {}
            for index, cell in enumerate(sheet[row_number], start=1):
                canonical = HEADER_KEYS.get(_header_key(cell.value))
                if canonical:
                    mapping[canonical] = index
            if set(REQUIRED_COLUMNS).issubset(mapping):
                yield sheet, row_number, mapping
                break


def _parse_file(path: Path) -> tuple[list[dict[str, Any]], int, int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    records: list[dict[str, Any]] = []
    invalid_rows = 0
    period_mismatches = 0
    found_sheet = False
    try:
        for sheet, header_row, columns in _workplace_sheets(workbook):
            found_sheet = True
            for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
                row = {name: values[index - 1] if index <= len(values) else None for name, index in columns.items()}
                if not any(value not in (None, "") for value in row.values()):
                    continue
                invoice = _invoice_number(row["Invoice Number"])
                amount = _amount(row["Amount"])
                description = re.sub(r"\s+", " ", str(row["Description"] or "")).strip()
                invoice_date = _invoice_date(row["Date"], row["Period"])
                period_month = _period_month(row["Period"])
                if not invoice or amount is None or not description or invoice_date is None:
                    invalid_rows += 1
                    continue
                month = invoice_date.replace(day=1)
                if period_month and period_month != month:
                    period_mismatches += 1
                records.append(
                    {
                        "invoice_number": invoice,
                        "amount_inr": amount,
                        "description": description[:500],
                        "date": invoice_date.isoformat(),
                        "month": month.isoformat(),
                        "period": month.strftime("%b %Y"),
                        "fiscal_year": fiscal_year_for_month(month),
                        "cost_type": _cost_type(description),
                    }
                )
    finally:
        workbook.close()
    if not found_sheet:
        raise ValueError(f"No sheet contains the required columns: {', '.join(REQUIRED_COLUMNS)}")
    return records, invalid_rows, period_mismatches


def discover_workplace_files(base_dir: Path | None) -> list[Path]:
    workplace_dir = base_dir / "workplace" if base_dir else None
    if workplace_dir is None or not workplace_dir.is_dir():
        return []
    return sorted(
        (
            path
            for path in workplace_dir.rglob("*")
            if path.is_file()
            and not path.is_symlink()
            and not path.name.startswith("~$")
            and path.suffix.casefold() in SUPPORTED_WORKBOOK_SUFFIXES
            and not any(part.startswith(".") for part in path.relative_to(workplace_dir).parts)
        ),
        key=lambda path: path.relative_to(workplace_dir).as_posix().casefold(),
    )


def workplace_signature(base_dir: Path | None) -> tuple[tuple[str, int, int], ...]:
    signature = []
    for path in discover_workplace_files(base_dir):
        stat = path.stat()
        signature.append((str(path), stat.st_mtime_ns, stat.st_size))
    return tuple(signature)


def discover_contracts(base_dir: Path | None) -> list[dict[str, Any]]:
    contracts_dir = base_dir / "contracts" if base_dir else None
    if contracts_dir is None or not contracts_dir.is_dir():
        return []
    contracts = []
    for path in contracts_dir.rglob("*"):
        if (
            not path.is_file()
            or path.is_symlink()
            or path.name.startswith("~$")
            or any(part.startswith(".") for part in path.relative_to(contracts_dir).parts)
        ):
            continue
        relative = path.relative_to(contracts_dir).as_posix()
        stat = path.stat()
        contracts.append(
            {
                "id": relative,
                "name": path.name,
                "relative_path": relative,
                "icon": CONTRACT_ICON_TYPES.get(path.suffix.casefold(), "file"),
                "extension": path.suffix.casefold().lstrip(".") or "file",
                "size_bytes": stat.st_size,
                "modified_at": datetime.fromtimestamp(stat.st_mtime).astimezone().isoformat(timespec="seconds"),
            }
        )
    return sorted(contracts, key=lambda item: item["relative_path"].casefold())


def load_workplace(base_dir: Path | None) -> dict[str, Any]:
    files = discover_workplace_files(base_dir)
    contracts = discover_contracts(base_dir)
    by_invoice: dict[str, dict[str, Any]] = {}
    invalid_files = 0
    invalid_rows = 0
    period_mismatches = 0
    duplicate_count = 0
    conflicting_duplicates = 0
    parsed_files = 0

    for path in files:
        try:
            records, file_invalid_rows, file_period_mismatches = _parse_file(path)
        except Exception:
            invalid_files += 1
            continue
        parsed_files += 1
        invalid_rows += file_invalid_rows
        period_mismatches += file_period_mismatches
        for record in records:
            fingerprint = hashlib.sha256(
                "\x1f".join(str(record[key]) for key in ("invoice_number", "amount_inr", "description", "date")).encode("utf-8")
            ).hexdigest()
            record["fingerprint"] = fingerprint
            existing = by_invoice.get(record["invoice_number"])
            if existing is None:
                by_invoice[record["invoice_number"]] = record
                continue
            duplicate_count += 1
            if existing["fingerprint"] != fingerprint:
                conflicting_duplicates += 1

    records = sorted(by_invoice.values(), key=lambda row: (row["date"], row["invoice_number"]))
    for record in records:
        record.pop("fingerprint", None)
    months = sorted({record["month"] for record in records})
    fiscal_years = []
    by_fy: dict[str, set[str]] = defaultdict(set)
    for record in records:
        by_fy[record["fiscal_year"]].add(record["month"])
    for label in sorted(by_fy):
        start_year = 2000 + int(label[-2:])
        fiscal_years.append(
            {
                "label": label,
                "start": date(start_year, 5, 1).isoformat(),
                "end": date(start_year + 1, 4, 30).isoformat(),
                "loaded_months": len(by_fy[label]),
                "complete": len(by_fy[label]) == 12,
            }
        )

    total = sum(record["amount_inr"] for record in records)
    average_monthly = total / len(months) if months else 0.0
    historical_annualized = average_monthly * 12
    latest_fy = fiscal_years[-1]["label"] if fiscal_years else None
    next_fy = f"FY{(int(latest_fy[-2:]) + 1) % 100:02d}" if latest_fy else None
    forecast = {
        "increase_rate": 5.0,
        "average_monthly": round(average_monthly, 2),
        "historical_annualized": round(historical_annualized, 2),
        "expected_monthly": round(average_monthly * 1.05, 2),
        "expected_annual": round(historical_annualized * 1.05, 2),
        "fiscal_year": next_fy,
        "months_observed": len(months),
    }
    return {
        "records": records,
        "contracts": contracts,
        "forecast": forecast,
        "meta": {
            "months": months,
            "fiscal_years": fiscal_years,
            "file_count": parsed_files,
            "record_count": len(records),
            "duplicate_count": duplicate_count,
            "conflicting_duplicates": conflicting_duplicates,
            "invalid_files": invalid_files,
            "invalid_rows": invalid_rows,
            "period_mismatches": period_mismatches,
            "contract_count": len(contracts),
            "storage_mode": "local_filesystem",
        },
    }
