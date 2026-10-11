from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime
import hashlib
from pathlib import Path
import re
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}


class SurveyFileError(ValueError):
    """A privacy-safe reason why a survey workbook could not be parsed."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code

COMMENT_HEADER_TERMS = (
    "comment",
    "suggestion",
    "feedback",
    "anything else",
    "what can we improve",
    "area for improvement",
    "areas for improvement",
)
IDENTITY_HEADER_TERMS = (
    "email",
    "e-mail",
    "name",
    "employee id",
    "employee number",
    "respondent",
    "phone",
    "mobile",
)
EMAIL_PATTERN = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.IGNORECASE)
PHONE_PATTERN = re.compile(r"(?<!\d)(?:\+?\d[\s().-]*){9,14}(?!\d)")
EXPLICIT_ID_PATTERN = re.compile(
    r"\b(?:employee|emp|staff|respondent)\s*(?:id|number|no\.?|#)\s*[:#=-]?\s*[A-Z0-9-]+\b",
    re.IGNORECASE,
)
PERSON_NAME_PAIR_PATTERN = re.compile(
    r"\b[A-Z][A-Za-z'-]{1,40}\s*(?:&|and)\s*[A-Z][A-Za-z'-]{1,40}\b"
)
PERSON_NAME_CONTEXT_PATTERN = re.compile(
    r"\b(?i:thanks?|thank you|kudos)\s+(?:(?i:to)\s+)?[A-Z][A-Za-z'-]{1,40}\b",
)
HONORIFIC_NAME_PATTERN = re.compile(r"\b(?:Mr|Mrs|Ms|Miss|Dr)\.?\s+[A-Z][A-Za-z'-]{1,40}\b")

QUESTION_PATTERNS: dict[str, tuple[str, ...]] = {
    "food": ("food and fresh beverage", "quality of lunch in the pantry"),
    "enhancements": ("recent workplace enhancements", "recent upgrades to workplace amenities"),
    "work_areas": ("work areas",),
    "meeting_rooms": ("meeting rooms",),
    "pantry": ("pantry",),
    "washroom": ("washroom",),
    "environment": ("air quality", "lighting", "temperature", "noise"),
    "reception": ("reception area",),
    "kitchen_breakout": ("kitchen", "breakout area"),
    "overall": ("overall experience with sequoia india facilities",),
}

TOUCHPOINT_LABELS = {
    "work_areas": "Work Areas",
    "meeting_rooms": "Meeting Rooms",
    "pantry": "Pantry",
    "washroom": "Washroom",
    "environment": "Air Quality, Lighting, Temperature & Noise",
    "reception": "Reception Area",
    "kitchen_breakout": "Kitchen/Breakout Area",
}

# Survey identity fields are deliberately outside the extraction allowlist.
# Fixed analytical columns and narrowly recognised comment-question columns are
# requested independently; Email, Name, ID and other response columns are never
# selected from a response row.
ALLOWED_RESPONSE_FIELDS = (
    "start_time",
    "completion_time",
    "food",
    "enhancements",
    "overall",
    *TOUCHPOINT_LABELS,
)


def _normalized(value: Any) -> str:
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip().casefold()


def _datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            converted = from_excel(value)
        except (TypeError, ValueError, OverflowError):
            return None
        if isinstance(converted, datetime):
            return converted
        if isinstance(converted, date):
            return datetime.combine(converted, datetime.min.time())
        return None
    raw = str(value or "").strip()
    if not raw:
        return None
    try:
        return datetime.fromisoformat(raw)
    except ValueError:
        pass
    # Microsoft Forms exports use month/day slash dates such as
    # 10/6/26 for 6 October 2026. Day/month is also accepted when the value is
    # unambiguous, while native Excel date cells avoid this ambiguity entirely.
    for fmt in (
        "%m/%d/%Y %H:%M:%S",
        "%m/%d/%Y %H:%M",
        "%m/%d/%Y %I:%M:%S %p",
        "%m/%d/%Y %I:%M %p",
        "%m/%d/%Y",
        "%m/%d/%y %H:%M:%S",
        "%m/%d/%y %H:%M",
        "%m/%d/%y %I:%M:%S %p",
        "%m/%d/%y %I:%M %p",
        "%m/%d/%y",
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%d/%m/%Y %I:%M:%S %p",
        "%d/%m/%Y %I:%M %p",
        "%d/%m/%Y",
        "%d/%m/%y %H:%M:%S",
        "%d/%m/%y %H:%M",
        "%d/%m/%y %I:%M:%S %p",
        "%d/%m/%y %I:%M %p",
        "%d/%m/%y",
    ):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            pass
    return None


def _rating(value: Any) -> int | None:
    try:
        number = int(float(value))
    except (TypeError, ValueError):
        return None
    return number if 1 <= number <= 5 else None


def _is_satisfied(value: Any) -> bool | None:
    normalized = _normalized(value)
    if normalized in {"satisfied", "very satisfied"}:
        return True
    if normalized in {"dissatisfied", "very dissatisfied"}:
        return False
    return None


def _satisfaction_label(value: Any) -> str | None:
    normalized = _normalized(value)
    labels = {
        "very dissatisfied": "Very dissatisfied",
        "dissatisfied": "Dissatisfied",
        "neutral": "Neutral",
        "neither satisfied nor dissatisfied": "Neutral",
        "satisfied": "Satisfied",
        "very satisfied": "Very satisfied",
    }
    return labels.get(normalized)


def _header_map(values: Iterable[Any]) -> dict[str, int]:
    headers = [(index, _normalized(value)) for index, value in enumerate(values)]
    mapping: dict[str, int] = {}
    for field, patterns in QUESTION_PATTERNS.items():
        for index, header in headers:
            if field in {"food", "enhancements"}:
                matched = any(pattern in header for pattern in patterns)
            elif field == "meeting_rooms":
                matched = header == "meeting rooms"
            elif field in {"work_areas", "pantry", "washroom"}:
                matched = header == patterns[0]
            else:
                matched = all(pattern in header for pattern in patterns)
            if matched:
                mapping[field] = index
                break
    for index, header in headers:
        if header == "start time":
            mapping["start_time"] = index
        elif header == "completion time":
            mapping["completion_time"] = index
    return mapping


def _comment_header_map(values: Iterable[Any], analytical_columns: set[int]) -> list[tuple[int, str]]:
    comments: list[tuple[int, str]] = []
    for index, value in enumerate(values):
        header = re.sub(r"\s+", " ", str(value or "")).strip()
        normalized = header.casefold()
        if not normalized or index in analytical_columns:
            continue
        if any(term in normalized for term in IDENTITY_HEADER_TERMS):
            continue
        if not any(term in normalized for term in COMMENT_HEADER_TERMS):
            continue
        if "suggestion" in normalized or "feedback" in normalized:
            topic = "Suggestions & feedback"
        elif "improve" in normalized:
            topic = "Improvement ideas"
        else:
            topic = "Comments"
        comments.append((index, topic))
    return comments


def _safe_comment(value: Any, limit: int = 2000) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()[:limit]
    if not text:
        return ""
    text = EMAIL_PATTERN.sub("[email removed]", text)
    text = PHONE_PATTERN.sub("[phone removed]", text)
    text = EXPLICIT_ID_PATTERN.sub("[employee identifier removed]", text)
    text = PERSON_NAME_PAIR_PATTERN.sub("[names removed]", text)
    text = PERSON_NAME_CONTEXT_PATTERN.sub("[name removed]", text)
    return HONORIFIC_NAME_PATTERN.sub("[name removed]", text)


def _survey_sheets(workbook) -> Iterable[tuple[Any, int, dict[str, int], list[tuple[int, str]]]]:
    # Food and workplace-enhancement questions were added or reworded across
    # survey versions. Start time and overall experience are the stable core;
    # every other available rating is included dynamically.
    required = {"start_time", "overall"}
    for sheet in workbook.worksheets:
        for row_number in range(1, min(sheet.max_row, 10) + 1):
            values = [sheet.cell(row_number, column).value for column in range(1, sheet.max_column + 1)]
            mapping = _header_map(values)
            if required.issubset(mapping):
                yield sheet, row_number, mapping, _comment_header_map(values, set(mapping.values()))
                break


def _response_fingerprint(record: dict[str, Any]) -> str:
    values = [
        record.get("started_at"),
        record.get("completed_at"),
        record.get("food"),
        record.get("food_response"),
        record.get("enhancements"),
        record.get("overall"),
        *(record.get("touchpoints", {}).get(key) for key in TOUCHPOINT_LABELS),
        *(record.get("touchpoint_responses", {}).get(key) for key in TOUCHPOINT_LABELS),
        *(comment.get("text") for comment in record.get("comments", [])),
    ]
    return hashlib.sha256("\x1f".join(str(value) for value in values).encode("utf-8")).hexdigest()


def _parse_file(path: Path) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise SurveyFileError(
            "unreadable_workbook",
            "The workbook could not be opened. It may be encrypted, incomplete, or damaged.",
        ) from exc
    records: list[dict[str, Any]] = []
    recognised_sheets = 0
    try:
        # Some Microsoft Forms exports contain complete worksheet XML but an
        # incorrect declared dimension of A1. Excel repairs this while opening
        # the file; openpyxl read-only mode needs the dimension reset explicitly.
        for sheet in workbook.worksheets:
            if sheet.max_row == 1 and sheet.max_column == 1 and hasattr(sheet, "reset_dimensions"):
                sheet.reset_dimensions()
                sheet.calculate_dimension(force=True)
        for sheet, header_row, columns, comment_columns in _survey_sheets(workbook):
            recognised_sheets += 1
            selected_fields = [field for field in ALLOWED_RESPONSE_FIELDS if field in columns]
            analytical_iterators = [
                sheet.iter_rows(
                    min_row=header_row + 1,
                    min_col=columns[field] + 1,
                    max_col=columns[field] + 1,
                    values_only=True,
                )
                for field in selected_fields
            ]
            comment_iterators = [
                sheet.iter_rows(
                    min_row=header_row + 1,
                    min_col=column + 1,
                    max_col=column + 1,
                    values_only=True,
                )
                for column, _topic in comment_columns
            ]
            # Each iterator exposes exactly one approved analytical or comment
            # column. Identity columns are never materialised as application values.
            for selected_cells in zip(*analytical_iterators, *comment_iterators):
                analytical_cells = selected_cells[: len(selected_fields)]
                comment_cells = selected_cells[len(selected_fields) :]
                values = {field: cells[0] for field, cells in zip(selected_fields, analytical_cells)}
                started = _datetime(values.get("start_time"))
                if started is None:
                    continue
                comments = [
                    {"topic": topic, "text": text}
                    for (_column, topic), cells in zip(comment_columns, comment_cells)
                    if (text := _safe_comment(cells[0]))
                ]
                touchpoint_responses = {
                    field: _satisfaction_label(values.get(field)) for field in TOUCHPOINT_LABELS
                }
                touchpoints = {field: _is_satisfied(values.get(field)) for field in TOUCHPOINT_LABELS}
                records.append(
                    {
                        "started_at": started.isoformat(),
                        "completed_at": (_datetime(values.get("completion_time")) or started).isoformat(),
                        "food": _is_satisfied(values.get("food")),
                        "food_response": _satisfaction_label(values.get("food")),
                        "enhancements": _rating(values.get("enhancements")),
                        "overall": _rating(values.get("overall")),
                        "touchpoints": touchpoints,
                        "touchpoint_responses": touchpoint_responses,
                        "comments": comments,
                    }
                )
    except Exception as exc:
        raise SurveyFileError(
            "processing_error",
            "The workbook opened, but its survey rows could not be processed.",
        ) from exc
    finally:
        workbook.close()
    if not recognised_sheets:
        raise SurveyFileError(
            "unrecognised_headers",
            "The required survey headers were not found together within the first ten rows of a worksheet.",
        )
    if not records:
        raise SurveyFileError(
            "no_valid_start_times",
            "Survey headers were found, but no valid Start time values were available.",
        )
    survey_start = min(datetime.fromisoformat(record["started_at"]) for record in records)
    survey_month = survey_start.date().replace(day=1).isoformat()
    for record in records:
        record["survey_month"] = survey_month
        record["fingerprint"] = _response_fingerprint(record)
    return records


def discover_survey_files(base_dir: Path | None) -> list[Path]:
    surveys_dir = base_dir / "surveys" if base_dir else None
    if surveys_dir is None or not surveys_dir.is_dir():
        return []
    return sorted(
        (
            path
            for path in surveys_dir.rglob("*")
            if path.is_file()
            and not path.is_symlink()
            and not path.name.startswith("~$")
            and path.suffix.casefold() in SUPPORTED_SUFFIXES
            and not any(part.startswith(".") for part in path.relative_to(surveys_dir).parts)
        ),
        key=lambda path: path.relative_to(surveys_dir).as_posix().casefold(),
    )


def survey_signature(base_dir: Path | None) -> tuple[tuple[str, int, int], ...]:
    signature = []
    for path in discover_survey_files(base_dir):
        stat = path.stat()
        signature.append((str(path), stat.st_mtime_ns, stat.st_size))
    return tuple(signature)


def _percentage(numerator: int, denominator: int) -> float | None:
    return round(numerator / denominator * 100, 1) if denominator else None


def _average(values: list[int]) -> float | None:
    return round(sum(values) / len(values), 2) if values else None


def load_surveys(base_dir: Path | None) -> dict[str, Any]:
    files = discover_survey_files(base_dir)
    if not files:
        return {
            "periods": [],
            "meta": {
                "months": [],
                "file_count": 0,
                "discovered_files": 0,
                "response_count": 0,
                "duplicate_count": 0,
                "invalid_files": 0,
                "invalid_reasons": {},
            },
        }

    by_fingerprint: dict[str, dict[str, Any]] = {}
    invalid_files = 0
    invalid_reasons: Counter[str] = Counter()
    parsed_count = 0
    raw_response_count = 0
    for path in files:
        try:
            records = _parse_file(path)
        except SurveyFileError as exc:
            invalid_files += 1
            invalid_reasons[exc.code] += 1
            continue
        except Exception:
            invalid_files += 1
            invalid_reasons["processing_error"] += 1
            continue
        parsed_count += 1
        raw_response_count += len(records)
        for record in records:
            by_fingerprint.setdefault(record["fingerprint"], record)

    records = list(by_fingerprint.values())
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        grouped[record["survey_month"]].append(record)

    periods = []
    for month in sorted(grouped):
        rows = grouped[month]
        overall = [row["overall"] for row in rows if row["overall"] is not None]
        enhancements = [row["enhancements"] for row in rows if row["enhancements"] is not None]
        food = [row["food"] for row in rows if row["food"] is not None]
        food_responses = [row["food_response"] for row in rows if row.get("food_response") is not None]
        touchpoints: dict[str, dict[str, Any]] = {}
        all_touchpoint_values: list[bool] = []
        for field, label in TOUCHPOINT_LABELS.items():
            values = [row["touchpoints"][field] for row in rows if row["touchpoints"].get(field) is not None]
            responses = [
                row["touchpoint_responses"][field]
                for row in rows
                if row.get("touchpoint_responses", {}).get(field) is not None
            ]
            satisfied = sum(bool(value) for value in values)
            all_touchpoint_values.extend(values)
            touchpoints[field] = {
                "label": label,
                "satisfied": satisfied,
                "responses": len(values),
                "rate": _percentage(satisfied, len(values)),
                "distribution": dict(Counter(responses)),
            }
        touchpoint_satisfied = sum(bool(value) for value in all_touchpoint_values)
        comments = [comment for row in rows for comment in row.get("comments", [])]
        commenting_responses = sum(bool(row.get("comments")) for row in rows)
        periods.append(
            {
                "month": month,
                "response_count": len(rows),
                "overall_average": _average(overall),
                "overall_favorable_rate": _percentage(sum(value >= 4 for value in overall), len(overall)),
                "enhancement_average": _average(enhancements),
                "enhancement_favorable_rate": _percentage(sum(value >= 4 for value in enhancements), len(enhancements)),
                "food_satisfaction_rate": _percentage(sum(bool(value) for value in food), len(food)),
                "touchpoint_satisfaction_rate": _percentage(touchpoint_satisfied, len(all_touchpoint_values)),
                "comment_count": len(comments),
                "commenting_response_count": commenting_responses,
                "comment_participation_rate": _percentage(commenting_responses, len(rows)),
                "comments": comments,
                "touchpoints": touchpoints,
                "rating_distributions": {
                    "overall": {str(score): Counter(overall).get(score, 0) for score in range(1, 6)},
                    "enhancements": {str(score): Counter(enhancements).get(score, 0) for score in range(1, 6)},
                },
                "satisfaction_distributions": {
                    "food": dict(Counter(food_responses)),
                },
            }
        )

    return {
        "periods": periods,
        "meta": {
            "months": [period["month"] for period in periods],
            "file_count": parsed_count,
            "discovered_files": len(files),
            "response_count": len(records),
            "duplicate_count": max(0, raw_response_count - len(records)),
            "invalid_files": invalid_files,
            "invalid_reasons": dict(invalid_reasons),
            "storage_mode": "local_with_masked_comments",
        },
    }
